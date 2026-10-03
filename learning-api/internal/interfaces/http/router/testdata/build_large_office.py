"""Build optional near-limit Office fixtures outside the repository.

Run with the bundled Python runtime and an output directory, then set
STARLINE_LARGE_OFFICE_FIXTURE_DIR when running the corresponding Go test.
"""

import sys
from pathlib import Path
from zipfile import ZIP_STORED, ZipFile

from docx import Document
from docx.shared import Inches as WordInches
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches as SlideInches


def store_package(source, target):
    with ZipFile(source) as original, ZipFile(target, "w", compression=ZIP_STORED) as output:
        for name in original.namelist():
            output.writestr(name, original.read(name))
    size = target.stat().st_size
    if not 48 * 1024 * 1024 < size < 50 * 1024 * 1024:
        raise ValueError(f"Fixture outside the intended upload range: {target.name}: {size}")
    print(f"{target.name}: {size} bytes")


def main():
    directory = Path(sys.argv[1])
    directory.mkdir(parents=True, exist_ok=True)
    picture = directory / "large-image.bmp"
    image = Image.new("RGB", (4096, 4160), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=130)
    for row in range(16):
        for col in range(16):
            draw.rectangle((col * 256, row * 260, (col + 1) * 256 - 1, (row + 1) * 260 - 1), fill=(col * 16, row * 16, 128))
    draw.text((120, 120), "Large teaching plan image", fill="white", font=font)
    draw.text((120, 3870), "Bottom content remains visible", fill="black", font=font)
    image.save(picture, format="BMP")

    document = Document()
    document.add_heading("Large Word teaching plan", 0)
    document.add_paragraph("First page: lesson preparation and image content.")
    document.add_picture(str(picture), width=WordInches(4.5))
    document.add_page_break()
    document.add_heading("Second Word page", 0)
    document.add_paragraph("Review questions and teacher notes remain on the second page.")
    small = directory / "compressed.docx"
    document.save(small)
    store_package(small, directory / "large-teaching-plan.docx")

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[5])
    first.shapes.title.text = "Large PPT teaching plan"
    first.shapes.add_picture(str(picture), SlideInches(0.7), SlideInches(1.4), width=SlideInches(5.2))
    second = deck.slides.add_slide(deck.slide_layouts[1])
    second.shapes.title.text = "Second PPT slide"
    second.placeholders[1].text = "Review questions\nTeacher preparation notes"
    small = directory / "compressed.pptx"
    deck.save(small)
    store_package(small, directory / "large-teaching-plan.pptx")


if __name__ == "__main__":
    main()
